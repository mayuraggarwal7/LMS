"""Import and export of institutional lesson plans.

Import: many colleges already keep a Word lesson plan with a table like

    Sr. No. | Proposed Date | Topics | Delivery Mode | CO | Actual Date | Remark
    3 | 22/7 (1) 26/7(2) | 2.1 Supply Chain Performance: Bullwhip effect ... | Lecture | CO2 | 26/7 | 9hrs
      | 15/8 | Independence Day
      | 12/9 | Mid Semester Examination (MSE)

`parse_lesson_plan` reads that table (from .doc/.docx/.pdf text produced by
`syllabus.extract_text`) and `apply_lesson_plan` turns it into units, topics
and plan items (one per lecture hour, "(2)" = two lectures that day),
plus holidays.

Export: `lesson_plan_rows` gives the same table for a course (proposed date
from the synced timetable, actual date and remark from delivered sessions)
and `build_docx` renders it as a Word document in the same layout.
"""

import io
import math
import re
from datetime import date

from django.db import transaction

DATE_RE = re.compile(r"(\d{1,2})\s*[/.-]\s*(\d{1,2})(?:\s*[/.-]\s*(\d{2,4}))?\s*(?:\(\s*(\d+(?:\.\d+)?)\s*\))?")
HOURS_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(?:hrs?|hours?|lec(?:ture)?s?)\b", re.I)
CO_CODE_RE = re.compile(r"\bC\.?O\.?\s*-?\s*(\d{1,2})\b", re.I)
SUBSECTION_RE = re.compile(r"(?:(?<=\s)|^)(\d{1,2})\.(\d{1,2})\s+(?=[A-Za-z(])")
EXAM_RE = re.compile(r"(mid[\s-]*sem|\bmse\b|end[\s-]*sem|\bese\b|\bexam(ination)?\b|\btest\b)", re.I)

HEADER_KEYS = [
    ("sr", r"^(sr|s\.?\s*no|sl|no\.?$|#)"),
    ("proposed", r"(proposed|planned|plan(ned)? date|tentative|^date)"),
    ("topics", r"(topic|content|syllabus|details)"),
    ("mode", r"(delivery|mode|method|pedagogy)"),
    ("co", r"^(co|cos|course outcome|c\.o)"),
    ("actual", r"(actual|conducted|completed on)"),
    ("remark", r"(remark|status|hours|hrs)"),
]


def academic_start_year(text, fallback=None):
    m = re.search(r"(20\d{2})\s*[-–/]\s*(20)?(\d{2})", text or "")
    if m:
        return int(m.group(1))
    return fallback


def resolve_date(day, month, year, start_year):
    """dd/mm with an optional year; academic years start around June."""
    day, month = int(day), int(month)
    if year:
        year = int(year)
        year = year + 2000 if year < 100 else year
    elif start_year:
        year = start_year if month >= 6 else start_year + 1
    else:
        year = date.today().year
    try:
        return date(year, month, day)
    except ValueError:
        return None


def parse_dates(text, start_year):
    """[(date, lectures)] from '12/7 (2) 15/7 (1)'."""
    result = []
    for d, m, y, n in DATE_RE.findall(text or ""):
        when = resolve_date(d, m, y, start_year)
        if when:
            result.append((when, float(n) if n else 1.0))
    return result


def _map_header(cells):
    mapping = {}
    for idx, cell in enumerate(cells):
        label = cell.strip().lower()
        for key, pattern in HEADER_KEYS:
            if key not in mapping and re.search(pattern, label):
                mapping[key] = idx
                break
    return mapping


def split_subsections(text):
    """'4.2 Warehousing: ... 4.3 Reverse logistics: ...' -> [(module, text), ...]."""
    marks = list(SUBSECTION_RE.finditer(text))
    if not marks:
        return [(None, text.strip())]
    parts = []
    if marks[0].start() > 0 and text[: marks[0].start()].strip():
        parts.append((None, text[: marks[0].start()].strip()))
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        parts.append((int(m.group(1)), text[m.end(): end].strip()))
    return parts


def parse_lesson_plan(text, start_year=None):
    """Return {"rows": [...], "events": [...]} or None if no lesson-plan table is found."""
    lines = (text or "").splitlines()
    start_year = start_year or academic_start_year(text)
    header = None
    rows, events = [], []
    for line in lines:
        if "|" not in line:
            continue
        cells = [c.strip() for c in line.split("|")]
        while cells and not cells[0]:
            cells.pop(0)
        if not cells:
            continue
        if header is None:
            mapping = _map_header(cells)
            if "topics" in mapping and ("proposed" in mapping or "sr" in mapping):
                header = mapping
            continue
        is_numbered = bool(re.fullmatch(r"\d{1,3}\.?", cells[0]))
        if not is_numbered and len(cells) <= 3:
            # special row: "15/8 | Independence Day" or "12/9 | Mid Semester Examination (MSE)"
            dates = parse_dates(cells[0], start_year)
            name = next((c for c in cells[1:] if c), "")
            if dates and name:
                kind = "exam" if EXAM_RE.search(name) else "holiday"
                exam = ""
                if kind == "exam":
                    exam = "MSE" if re.search(r"mid[\s-]*sem|\bmse\b", name, re.I) else (
                        "ESE" if re.search(r"end[\s-]*sem|\bese\b", name, re.I) else "")
                events.append({"date": dates[0][0].isoformat(), "name": name[:120], "kind": kind, "exam": exam})
            continue
        if not is_numbered:
            offset = -1 if "sr" in header else 0
        else:
            offset = 0

        def cell(key):
            idx = header.get(key)
            if idx is None:
                return ""
            idx += offset
            return cells[idx] if 0 <= idx < len(cells) else ""

        topics = cell("topics")
        if not topics:
            continue
        dates = parse_dates(cell("proposed"), start_year)
        remark = " ".join([cell("remark"), cell("actual")])
        hours_match = HOURS_RE.search(remark)
        cos = sorted({int(n) for n in CO_CODE_RE.findall(cell("co"))})
        rows.append({
            "sr": cells[0] if is_numbered else "",
            "dates": [(d.isoformat(), n) for d, n in dates],
            "topics": re.sub(r"\s+", " ", topics).strip(),
            "mode": cell("mode") or "Lecture",
            "cos": [f"CO{n}" for n in cos],
            "hours": float(hours_match.group(1)) if hours_match else None,
            "not_covered": bool(re.search(r"not\s+covered", remark, re.I)),
        })
    if not rows:
        return None
    return {"rows": rows, "events": events, "start_year": start_year}


def lecture_count(row):
    if row["dates"]:
        return max(1, math.ceil(sum(n for _, n in row["dates"])))
    if row.get("hours"):
        return max(1, math.ceil(row["hours"]))
    return 1


def row_segments(row):
    """[(module or None, lead title or None, [topic names])] for one lesson-plan row."""
    from .syllabus import split_topics

    segments = []
    for module, text in split_subsections(row["topics"]):
        lead = None
        if re.match(r"case\s+stud", text, re.I):
            names = [re.sub(r"\s+", " ", text).strip(" .")]
        else:
            names = split_topics(text)
            if names and ":" in names[0]:
                head, _, tail = names[0].partition(":")
                if len(head) < 70 and tail.strip():
                    lead = head.strip()
                    names[0] = tail.strip()[0].upper() + tail.strip()[1:]
        segments.append((module, lead, names))
    return segments


def derive_units(lesson_plan):
    """Units/topics from the lesson plan's 'm.n' numbering (for syllabi without unit headings)."""
    units, pending, pending_hours, extra = {}, [], 0, None
    for row in lesson_plan["rows"]:
        hours_assigned = False
        for module, lead, names in row_segments(row):
            if module is None:
                if not units:
                    pending += names
                    pending_hours += lecture_count(row)
                    continue
                if extra is None:
                    extra = max(units) + 1
                    title = "Case studies & applications" if names and re.match(r"case\s+stud", names[0], re.I) else "Additional topics"
                    units[extra] = {"number": extra, "title": title, "hours": 0, "topics": []}
                module = extra
            if module not in units:
                title = lead or (names[0] if names else f"Module {module}")
                units[module] = {"number": module, "title": title[:120], "hours": 0, "topics": []}
            units[module]["topics"] += names
            if not hours_assigned:
                units[module]["hours"] += lecture_count(row)
                hours_assigned = True
        if pending and units:
            first = units[min(units)]
            intro_title = pending[0]
            first["title"] = intro_title[:120]
            first["topics"] = pending + first["topics"]
            first["hours"] += pending_hours
            pending, pending_hours = [], 0
    for u in units.values():
        seen, topics = set(), []
        for t in u["topics"]:
            if t.lower() not in seen:
                seen.add(t.lower())
                topics.append(t[:300])
        u["topics"] = topics
    return [units[n] for n in sorted(units)]


@transaction.atomic
def apply_lesson_plan(course, lesson_plan, replace=True):
    """Create plan items (and holidays) from a parsed lesson plan. Returns number of items."""
    from delivery.models import Holiday

    from .models import CourseOutcome, PlanItem, Topic, Unit
    from .planner import _chunk, build_session_plan

    if replace:
        course.plan_items.all().delete()
    outcomes = list(CourseOutcome.objects.filter(course=course))
    by_code = {o.code.upper(): o for o in outcomes}
    units = {u.number: u for u in course.units.all()}
    if not units:
        for u in derive_units(lesson_plan):
            unit = Unit.objects.create(course=course, number=u["number"], title=u["title"], hours=u["hours"])
            for i, t in enumerate(u["topics"]):
                Topic.objects.create(unit=unit, title=t[:300], order=i)
            units[unit.number] = unit
    unit_list = [units[n] for n in sorted(units)]

    def in_range(d):
        return (not course.start_date or d >= course.start_date) and (not course.end_date or d <= course.end_date)

    seq = course.plan_items.count()
    created = 0
    current_unit = unit_list[0] if unit_list else None
    for row in lesson_plan["rows"]:
        segments = row_segments(row)
        modules = [m for m, _, _ in segments if m]
        if modules and modules[0] in units:
            current_unit = units[modules[0]]
        elif not modules and re.match(r"case\s+stud", row["topics"], re.I) and unit_list:
            current_unit = unit_list[-1]
        topic_names = [t for _, _, names in segments for t in names] or [row["topics"][:300]]
        topics = []
        for name in topic_names:
            match = None
            if current_unit:
                match = next((t for t in current_unit.topics.all() if t.title.lower() == name.lower()), None)
                if match is None:
                    match = Topic.objects.create(unit=current_unit, title=name[:300], order=current_unit.topics.count())
            if match:
                topics.append(match)
        n = lecture_count(row)
        dates = []
        for d, count in row["dates"]:
            dates += [date.fromisoformat(d)] * max(1, math.ceil(count))
        chunks = _chunk(topics, n) if topics else [[]] * n
        row_cos = [by_code[c] for c in row["cos"] if c in by_code]
        for i, chunk in enumerate(chunks):
            seq += 1
            if chunk and current_unit:
                data = build_session_plan(course, current_unit, chunk, i, len(chunks), outcomes,
                                          unit_list.index(current_unit), len(unit_list),
                                          (i + 1, len(chunks)) if len(chunks) > 1 and len(chunk) == 1 and len(topics) == 1 else None)
            else:
                data = {"title": row["topics"][:400], "outcomes": [], "objective": "", "bloom_level": 2}
            if row_cos:
                data["outcomes"] = row_cos
                data["bloom_level"] = min(data.get("bloom_level", 2), max(o.bloom_level for o in row_cos))
            mode = row["mode"].strip().lower()
            session_type = "lab" if re.search(r"lab|practical", mode) else "tutorial" if "tutorial" in mode else "lecture"
            planned = dates[i] if i < len(dates) and in_range(dates[i]) else None
            item = PlanItem.objects.create(
                course=course, sequence=seq, session_type=session_type, unit=current_unit if chunk else None,
                planned_date=planned, **{k: v for k, v in data.items() if k != "outcomes"},
            )
            if row["mode"] and row["mode"].lower() not in ("lecture", "lab", "tutorial"):
                item.pedagogy = row["mode"][:200]
                item.save(update_fields=["pedagogy"])
            item.topics.set(chunk)
            item.outcomes.set(data["outcomes"])
            if current_unit and data["outcomes"]:
                current_unit.outcomes.add(*data["outcomes"])
            created += 1

    for event in lesson_plan.get("events", []):
        when = date.fromisoformat(event["date"])
        if event["kind"] == "holiday" and in_range(when):
            Holiday.objects.get_or_create(course=course, date=when, defaults={"name": event["name"]})
    return created


# --- export ---------------------------------------------------------------------------

def lesson_plan_rows(course):
    """Chronological rows for the institutional lesson plan table."""
    from delivery.models import ClassSession

    today = date.today()
    rows = []
    sessions = {}
    for s in course.sessions.exclude(status=ClassSession.CANCELLED).filter(plan_item__isnull=False).order_by("date", "start_time"):
        sessions.setdefault(s.plan_item_id, s)
    for item in course.plan_items.select_related("unit").prefetch_related("outcomes"):
        s = sessions.get(item.pk)
        proposed = s.date if s else item.planned_date
        if s and s.status == ClassSession.COMPLETED:
            actual = s.date
            remark = "Partially covered" if s.coverage == "partial" else "Covered"
        else:
            actual = None
            remark = "Not covered" if proposed and proposed < today else ""
        rows.append({
            "kind": "item", "date": proposed, "item": item, "proposed": proposed, "actual": actual, "remark": remark,
            "topics": item.title, "unit": item.unit.number if item.unit else "", "mode": item.get_session_type_display(),
            "cos": ", ".join(o.code for o in item.outcomes.all()),
        })
    for h in course.holidays.all():
        rows.append({"kind": "event", "date": h.date, "topics": h.name or "Holiday"})
    for e in course.exams.all():
        if e.date:
            rows.append({"kind": "event", "date": e.date, "topics": e.title})
    far = date(9999, 1, 1)
    rows.sort(key=lambda r: (r["date"] or far, 0 if r["kind"] == "event" else 1, r["item"].sequence if r["kind"] == "item" else 0))
    sr = 0
    for r in rows:
        if r["kind"] == "item":
            sr += 1
            r["sr"] = sr
    return rows


def build_docx(course):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)

    def centered(text, bold=True, size=12):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)

    if course.department:
        centered(course.department, size=13)
    if course.academic_year:
        centered(f"(Academic Year: {course.academic_year})", bold=False, size=11)
    document.add_paragraph().add_run(f"Course Code: {course.code}").bold = True
    document.add_paragraph().add_run(f"Course Name: {course.title}").bold = True
    document.add_paragraph().add_run(f"Course Teacher: {course.teacher.display_name}").bold = True
    document.add_paragraph("Course Outcomes (CO): At the end of the course students will be able to")
    outcomes = list(course.outcomes.all())
    if outcomes:
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for o in outcomes:
            cells = table.add_row().cells
            cells[0].text = o.code
            cells[1].text = o.description

    centered("Course Lesson Plan", size=12)
    headers = ["Sr. No.", "Proposed Date", "Topics", "Delivery Mode", "CO", "Actual Date", "Remark"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for r in lesson_plan_rows(course):
        cells = table.add_row().cells
        fmt = lambda d: d.strftime("%d/%m/%Y") if d else ""  # noqa: E731
        if r["kind"] == "event":
            cells[1].text = fmt(r["date"])
            merged = cells[2].merge(cells[-1])
            merged.text = ""
            merged.paragraphs[0].add_run(r["topics"]).bold = True
            continue
        values = [str(r["sr"]), fmt(r["proposed"]), (f"{r['unit']}: " if r["unit"] else "") + r["topics"],
                  r["mode"], r["cos"], fmt(r["actual"]), r["remark"]]
        for cell, text in zip(cells, values):
            cell.text = text

    books = list(course.references.all())
    if books:
        document.add_paragraph().add_run("Text Books / References:").bold = True
        for i, ref in enumerate(books, 1):
            document.add_paragraph(f"{i}. {ref.text}")
    document.add_paragraph()
    document.add_paragraph().add_run(f"Course Instructor: {course.teacher.display_name}").bold = True
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
