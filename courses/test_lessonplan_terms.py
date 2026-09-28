import io
from datetime import date, time, timedelta
from decimal import Decimal

import docx
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import User
from classroom import grading
from classroom.models import Assignment
from courses import lessonplan, syllabus
from courses.models import AcademicTerm, Course, Enrollment
from courses.setup import build_course
from delivery.models import AttendanceRecord, ClassSession, TimetableSlot
from delivery.services import sync_timetable

# Same shape as a departmental lesson plan exported from Word (cells joined by " | ").
LESSON_PLAN_TEXT = """Mechanical Engineering
(Academic Year: 2024-2025)
Course Code: MEC702
Course Name: Logistics and Supply Chain Management
Course Outcomes (CO): At the End of the course students will be able to
CO.1 | Demonstrate a sound understanding of Logistics and Supply Chain Management concepts.
CO.2 | Identify the drivers of supply chain performance and risks in supply chain management.
CO.3 | Apply various techniques of inventory management.
Course Lesson Plan
Sr. No. | Proposed Date | Topics | Delivery Mode | CO | Actual Date | Remark
1 | 12/7 (2) 15/7 (1) | Objectives of a Supply Chain Management, Stages of Supply chain, Key issues in SCM | Lecture | CO1 | 12/7 15/7 | 3hrs
2 | 19/7 (2) | 1.1 Supplier Selection, Supplier quality audits, Contract management | Lecture | CO2 | 26/7
 | 3 | 22/7 (1) 26/7(2) | 2.1 Supply Chain Performance: Bullwhip effect and reduction, SCOR Model | Lecture | CO2 | 26/7 | 3hrs
15/8 | Independence Day
4 | 16/8(2) | 3.1 Inventory management: EOQ Model, ABC Analysis. 3.2 Replenishment systems | Lecture | CO3 | Not Covered
12/9 | Mid Semester Examination (MSE)
5 | 21/10(1) | Case Study – Jimny Maruti | Lecture | CO3 | 25/08 | 1hrs
21/11 | End Semester Examination (ESE)
Text Books:
1. R. K. Rajput "Supply Chain Management", S. Chand
Course Instructor: Prof. Example
"""


def teacher():
    return User.objects.create_user("t", password="pw123456", role=User.TEACHER, first_name="Tee")


class LessonPlanParserTests(TestCase):
    def test_parse_rows_dates_and_events(self):
        lp = lessonplan.parse_lesson_plan(LESSON_PLAN_TEXT)
        self.assertEqual(lp["start_year"], 2024)
        self.assertEqual(len(lp["rows"]), 5)
        first = lp["rows"][0]
        self.assertEqual(first["dates"], [("2024-07-12", 2.0), ("2024-07-15", 1.0)])
        self.assertEqual(first["cos"], ["CO1"])
        self.assertEqual(lessonplan.lecture_count(first), 3)
        self.assertTrue(lp["rows"][3]["not_covered"])
        events = {e["name"]: e for e in lp["events"]}
        self.assertEqual(events["Independence Day"]["kind"], "holiday")
        self.assertEqual(events["Independence Day"]["exam"], "")
        self.assertEqual(events["Mid Semester Examination (MSE)"]["exam"], "MSE")
        self.assertEqual(events["End Semester Examination (ESE)"]["date"], "2024-11-21")

    def test_derive_units_from_numbering(self):
        units = lessonplan.derive_units(lessonplan.parse_lesson_plan(LESSON_PLAN_TEXT))
        self.assertEqual([u["number"] for u in units], [1, 2, 3, 4])
        self.assertEqual(units[1]["title"], "Supply Chain Performance")
        self.assertEqual(units[2]["title"], "Inventory management")
        self.assertEqual(units[3]["title"], "Case studies & applications")
        self.assertEqual(sum(u["hours"] for u in units), 3 + 2 + 3 + 2 + 1)
        self.assertIn("Supplier Selection", units[0]["topics"])  # 1.x rows join the intro unit

    def test_syllabus_parser_stops_outcomes_at_table_and_keeps_initials(self):
        data = syllabus.parse_syllabus(LESSON_PLAN_TEXT)
        self.assertEqual(len(data["outcomes"]), 3)
        self.assertNotIn("Proposed", data["outcomes"][-1]["description"])
        self.assertEqual(data["textbooks"], ['R. K. Rajput "Supply Chain Management", S. Chand'])
        self.assertEqual(syllabus.guess_department(LESSON_PLAN_TEXT), "Mechanical Engineering")

    def test_docx_table_extraction(self):
        document = docx.Document()
        document.add_paragraph("Course Lesson Plan")
        table = document.add_table(rows=2, cols=5)
        for cell, text in zip(table.rows[0].cells, ["Sr. No.", "Proposed Date", "Topics", "Delivery Mode", "CO"]):
            cell.text = text
        for cell, text in zip(table.rows[1].cells, ["1", "5/8 (2)", "1.1 Sets, relations", "Lecture", "CO1"]):
            cell.text = text
        buf = io.BytesIO()
        document.save(buf)
        text = syllabus.extract_text(SimpleUploadedFile("lp.docx", buf.getvalue()))
        lp = lessonplan.parse_lesson_plan(text, start_year=2026)
        self.assertEqual(lp["rows"][0]["dates"], [("2026-08-05", 2.0)])


class LessonPlanCourseTests(TestCase):
    def setUp(self):
        self.teacher = teacher()
        self.lp = lessonplan.parse_lesson_plan(LESSON_PLAN_TEXT)
        data = syllabus.parse_syllabus(LESSON_PLAN_TEXT)
        data["units"] = lessonplan.derive_units(self.lp)
        info = {"code": "MEC702", "title": data["title"], "year": 4, "semester": 7,
                "start_date": date(2024, 7, 8), "end_date": date(2024, 11, 15)}
        self.course, self.summary = build_course(self.teacher, data, info, lesson_plan=self.lp)

    def test_plan_follows_lesson_plan(self):
        self.assertTrue(self.summary["from_lesson_plan"])
        items = list(self.course.plan_items.all())
        self.assertEqual(len(items), 11)  # one per lecture
        self.assertEqual(items[0].planned_date, date(2024, 7, 12))
        self.assertEqual(items[1].planned_date, date(2024, 7, 12))  # "(2)" = two lectures that day
        self.assertEqual(items[2].planned_date, date(2024, 7, 15))
        self.assertEqual([o.code for o in items[3].outcomes.all()], ["CO2"])
        self.assertEqual(items[-1].unit.title, "Case studies & applications")
        self.assertTrue(self.course.holidays.filter(date=date(2024, 8, 15)).exists())
        self.assertEqual(self.course.exams.get(component__kind="MSE").date, date(2024, 9, 12))

    def test_export_rows_and_docx(self):
        TimetableSlot.objects.create(course=self.course, weekday=0, start_time=time(10), end_time=time(11))
        sync_timetable(self.course)
        first = self.course.sessions.order_by("date").first()
        first.status = ClassSession.COMPLETED
        first.save()
        rows = lessonplan.lesson_plan_rows(self.course)
        items = [r for r in rows if r["kind"] == "item"]
        self.assertEqual(items[0]["remark"], "Covered")
        self.assertEqual(items[0]["actual"], first.date)
        self.assertEqual(items[1]["remark"], "Not covered")  # 2024 date already past
        self.assertTrue(any(r["kind"] == "event" and r["topics"] == "Independence Day" for r in rows))
        document = docx.Document(io.BytesIO(lessonplan.build_docx(self.course)))
        text = "\n".join(c.text for t in document.tables for row in t.rows for c in row.cells)
        self.assertIn("Proposed Date", text)
        self.assertIn("Independence Day", text)

    def test_import_view_replaces_plan(self):
        self.client.login(username="t", password="pw123456")
        upload = SimpleUploadedFile("plan.txt", LESSON_PLAN_TEXT.encode())
        r = self.client.post(reverse("plan_import", args=[self.course.pk]), {"file": upload, "keep_units": "on"})
        self.assertRedirects(r, reverse("course_plan", args=[self.course.pk]))
        self.assertEqual(self.course.plan_items.count(), 11)
        r = self.client.get(reverse("lesson_plan_document", args=[self.course.pk]) + "?format=docx")
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


class TermTests(TestCase):
    def setUp(self):
        self.teacher = teacher()
        self.odd = AcademicTerm.objects.create(name="Odd Term 2025", kind="odd", academic_year=2025)
        self.even = AcademicTerm.objects.create(name="Even Term 2026", kind="even", academic_year=2026, is_active=True)
        self.old = Course.objects.create(teacher=self.teacher, code="OLD1", title="Old course", term=self.odd)
        self.new = Course.objects.create(teacher=self.teacher, code="NEW1", title="New course", term=self.even)

    def test_single_active_term(self):
        self.odd.is_active = True
        self.odd.save()
        self.even.refresh_from_db()
        self.assertFalse(self.even.is_active)

    def test_dashboard_shows_active_term_then_switched_term(self):
        self.client.login(username="t", password="pw123456")
        r = self.client.get(reverse("dashboard"))
        self.assertContains(r, "NEW1")
        self.assertNotContains(r, "OLD1")
        self.client.post(reverse("term_switch"), {"term": self.odd.pk})
        r = self.client.get(reverse("dashboard"))
        self.assertContains(r, "OLD1")
        self.assertNotContains(r, "NEW1")

    def test_new_course_goes_into_viewed_term(self):
        self.client.login(username="t", password="pw123456")
        self.client.post(reverse("term_switch"), {"term": self.odd.pk})
        self.client.post(reverse("course_new"), {"syllabus_text": LESSON_PLAN_TEXT, "year": 4, "semester": 7})
        form = self.client.get(reverse("course_review")).context["form"]
        post = {k: ("" if v is None else v) for k, v in form.initial.items()}
        post.update({"opt_plan": "on", "opt_lesson_plan": "on"})
        self.client.post(reverse("course_review"), post)
        course = Course.objects.get(code="MEC702")
        self.assertEqual(course.term, self.odd)
        self.assertEqual(course.plan_items.count(), 11)


class PendingAttendanceAndCIETests(TestCase):
    def setUp(self):
        self.teacher = teacher()
        today = timezone.localdate()
        self.course = Course.objects.create(teacher=self.teacher, code="C1", title="Course",
                                            start_date=today - timedelta(days=14), end_date=today + timedelta(days=30))
        from courses.planner import generate_learning_plan
        from courses.models import Topic, Unit
        from exams.services import create_default_scheme

        unit = Unit.objects.create(course=self.course, number=1, title="Basics", hours=10)
        for i in range(10):
            Topic.objects.create(unit=unit, title=f"Topic {i}", order=i)
        generate_learning_plan(self.course)
        create_default_scheme(self.course)
        TimetableSlot.objects.create(course=self.course, weekday=today.weekday(), start_time=time(9), end_time=time(10))
        sync_timetable(self.course)
        self.students = []
        for i, (roll, prn, name) in enumerate([("2", "P9", "Zed"), ("1", "P5", "Amy")]):
            s = User.objects.create_user(f"s{i}", password="pw123456", roll_no=roll, prn=prn, first_name=name)
            Enrollment.objects.create(course=self.course, student=s)
            self.students.append(s)
        self.client.login(username="t", password="pw123456")

    def test_pending_list_and_bulk_delete_shifts_plan(self):
        r = self.client.get(reverse("pending_attendance"))
        pending = list(r.context["pending"])
        self.assertEqual(len(pending), 3)  # two past weeks + today
        first_item = pending[0].plan_item
        self.client.post(reverse("pending_attendance"), {"selected": [pending[0].pk]})
        pending[0].refresh_from_db()
        self.assertEqual(pending[0].status, ClassSession.CANCELLED)
        pending[1].refresh_from_db()
        self.assertEqual(pending[1].plan_item, first_item)  # plan moved to the next class
        # marking attendance removes a class from the pending list
        AttendanceRecord.objects.create(session=pending[1], student=self.students[0])
        self.assertEqual(len(self.client.get(reverse("pending_attendance")).context["pending"]), 1)

    def test_cie_report_ne_sorting_and_csv(self):
        a = Assignment.objects.create(course=self.course, title="A1", max_points=Decimal("10"), status=Assignment.PUBLISHED,
                                      component=self.course.components.get(name="ISE-1"))
        grading.set_total_score(grading.get_or_create_submission(a, self.students[0]), Decimal("8"))
        r = self.client.get(reverse("cie_report", args=[self.course.pk]) + "?sort=prn")
        rows = r.context["rows"]
        self.assertEqual([row["student"].prn for row in rows], ["P5", "P9"])
        self.assertIsNone(rows[0]["cie_total"])  # Amy: nothing entered -> NE
        self.assertEqual(rows[1]["cie_total"], Decimal("16.00"))  # 80% of the 20-mark ISE-1 bucket
        self.assertContains(r, "NE")
        csv = self.client.get(reverse("cie_report", args=[self.course.pk]) + "?format=csv").content.decode()
        self.assertIn("Total CIE", csv.splitlines()[0])
        self.assertIn("P9", csv)
