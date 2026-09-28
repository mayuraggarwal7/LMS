"""Syllabus extraction and parsing.

`extract_text` turns an uploaded PDF / DOCX / TXT into plain text.
`parse_syllabus` turns that text into a structured dict:

    {
      "code": "CS201", "title": "Data Structures", "credits": 4,
      "lecture_hours": 3, "tutorial_hours": 0, "practical_hours": 2,
      "outcomes": [{"code": "CO1", "description": "..."}],
      "units": [{"number": 1, "title": "...", "hours": 8, "topics": ["...", ...]}],
      "textbooks": ["..."], "references": ["..."], "experiments": ["..."],
    }

The heuristic parser handles the layouts used by most Indian engineering
universities (Unit/Module headings, "CO1:" outcome lists, "Text Books" and
"List of Experiments" sections). When ANTHROPIC_API_KEY is configured the
Claude-assisted parser in `courses.ai` is tried first and this is the fallback.

`to_editable` / `from_editable` convert the dict to and from the plain-text
format teachers review and correct on the setup screen.
"""

import io
import re

ROMAN = {"i": 1, "ii": 2, "iii": 3, "iv": 4, "v": 5, "vi": 6, "vii": 7, "viii": 8, "ix": 9, "x": 10, "xi": 11, "xii": 12}

UNIT_RE = re.compile(
    r"^\s*(?:unit|module|chapter|section)\s*[-:#.]?\s*(?P<num>[ivxlc]+|\d{1,2})\b\s*[:.\-–—)]*\s*(?P<rest>.*)$",
    re.IGNORECASE,
)
CO_RE = re.compile(r"^\s*(?:[-•*]\s*)?(?P<code>C\.?O\.?\s*-?\s*\d{1,2})\s*[:.\-–—)]*\s*(?P<text>.+)$", re.IGNORECASE)
HOURS_RE = re.compile(r"[\(\[]?\s*(\d{1,2})\s*(?:hours|hrs|hr|h|lectures|lecs|l)\b\.?\s*[\)\]]?", re.IGNORECASE)
TRAILING_NUM_RE = re.compile(r"\s+(\d{1,2})\s*$")
NUMBERED_RE = re.compile(r"^\s*(?:\[?\d{1,2}[\].)]|[-•*]|[a-z][.)])\s*(.+)$", re.IGNORECASE)

SECTION_PATTERNS = {
    "outcomes": re.compile(r"^\s*(course\s+outcomes?|learning\s+outcomes?|cos?\b|outcomes?)\s*[:\-]?", re.I),
    "objectives": re.compile(r"^\s*(course\s+objectives?|objectives?)\s*[:\-]?", re.I),
    "textbooks": re.compile(r"^\s*(text\s*books?|textbooks?|prescribed\s+books?)\s*[:\-]?", re.I),
    "references": re.compile(r"^\s*(reference\s*books?|references?|suggested\s+reading|further\s+reading)\s*[:\-]?", re.I),
    "online": re.compile(r"^\s*(online\s+resources?|web\s+resources?|e-?resources?|mooc)", re.I),
    "experiments": re.compile(r"^\s*(list\s+of\s+(?:experiments|practicals|lab(?:oratory)?\s+exercises)|laboratory\s+(?:work|experiments)|practicals?|lab\s+experiments?|suggested\s+experiments?)\s*[:\-]?", re.I),
    "assessment": re.compile(r"^\s*(assessment|evaluation\s+scheme|examination\s+scheme|scheme\s+of\s+evaluation|internal\s+assessment)", re.I),
    "contents": re.compile(r"^\s*(course\s+contents?|syllabus|detailed\s+syllabus|contents)\s*[:\-]?\s*$", re.I),
    "prerequisites": re.compile(r"^\s*(pre-?requisites?)\s*[:\-]?", re.I),
}


# --- text extraction ----------------------------------------------------------

def extract_text(uploaded_file):
    """Return plain text from an uploaded syllabus (PDF, DOCX, or text)."""
    name = (getattr(uploaded_file, "name", "") or "").lower()
    data = uploaded_file.read()
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if name.endswith(".docx"):
        import docx

        document = docx.Document(io.BytesIO(data))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    text = cell.text.strip()
                    if text and (not cells or cells[-1] != text):
                        cells.append(text)
                lines.append(" | ".join(cells))
        return "\n".join(lines)
    for encoding in ("utf-8", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""


# --- heuristics -----------------------------------------------------------------

def _to_int(token):
    token = token.strip().lower()
    if token.isdigit():
        return int(token)
    return ROMAN.get(token)


def _clean(line):
    return re.sub(r"\s+", " ", line.replace("\t", " ")).strip(" |:-–—.;")


def _section_of(line):
    stripped = line.strip()
    if len(stripped) > 60:
        return None
    for name, pattern in SECTION_PATTERNS.items():
        if pattern.match(stripped):
            # "CO1: ..." lines are outcomes, not section headers
            if name == "outcomes" and CO_RE.match(stripped):
                return None
            return name
    return None


def _extract_hours(text):
    """Return (text_without_hours, hours or None)."""
    match = HOURS_RE.search(text)
    if match:
        return _clean(text[: match.start()] + " " + text[match.end():]), int(match.group(1))
    trailing = TRAILING_NUM_RE.search(text)
    if trailing and len(text) > 6:
        return _clean(text[: trailing.start()]), int(trailing.group(1))
    return _clean(text), None


def split_topics(text):
    """Split a unit's content paragraph into individual topic strings."""
    text = re.sub(r"\s+", " ", text)
    # protect separators inside parentheses
    depth, buf, parts = 0, "", []
    for ch in text:
        if ch in "([":
            depth += 1
        elif ch in ")]" and depth:
            depth -= 1
        if depth == 0 and ch in ",;•":
            parts.append(buf)
            buf = ""
        elif depth == 0 and ch == ".":
            parts.append(buf)
            buf = ""
        else:
            buf += ch
    parts.append(buf)
    topics = []
    for part in parts:
        for piece in re.split(r"\s+[–—]\s+|\s-\s", part):
            piece = _clean(piece)
            piece = re.sub(r"^(and|&)\s+", "", piece, flags=re.I)
            if len(piece) >= 3 and not piece.isdigit():
                topics.append(piece[0].upper() + piece[1:])
    # de-duplicate, keep order
    seen, unique = set(), []
    for t in topics:
        key = t.lower()
        if key not in seen:
            seen.add(key)
            unique.append(t)
    return unique


def parse_syllabus(text):
    lines = [l.rstrip() for l in (text or "").replace("\r", "\n").split("\n")]
    result = {
        "code": "", "title": "", "credits": None,
        "lecture_hours": None, "tutorial_hours": None, "practical_hours": None,
        "prerequisites": "", "outcomes": [], "units": [],
        "textbooks": [], "references": [], "online": [], "experiments": [],
    }
    head = "\n".join(lines[:40])

    m = re.search(r"(?:course|subject|paper)\s*code\s*[:\-]?\s*([A-Z]{2,6}[\s-]?\d{2,4}[A-Z]?)", head, re.I)
    if not m:
        m = re.search(r"\b([A-Z]{2,5}[\s-]?\d{3,4}[A-Z]?)\b", head)
    if m:
        result["code"] = re.sub(r"\s", "", m.group(1)).upper()

    m = re.search(r"(?:course|subject|paper)\s*(?:title|name)\s*[:\-]\s*(.+)", head, re.I)
    if m:
        result["title"] = _clean(re.split(r"\s{2,}|\|", m.group(1))[0])
    else:
        for line in lines[:12]:
            cl = _clean(line)
            if cl and not re.search(r"code|credit|university|institute|college|semester|scheme|L\s*-?\s*T", cl, re.I) and len(cl) > 4 and not UNIT_RE.match(cl):
                result["title"] = re.sub(r"^[A-Z]{2,5}[\s-]?\d{3,4}[A-Z]?\s*[:\-]?\s*", "", cl)
                break

    m = re.search(r"credits?\s*[:\-]?\s*(\d+(?:\.\d)?)", head, re.I)
    if m:
        result["credits"] = float(m.group(1))

    m = re.search(r"\bL\s*[:\-]?\s*(\d)\s*[,/|-]?\s*T\s*[:\-]?\s*(\d)\s*[,/|-]?\s*P\s*[:\-]?\s*(\d)", head)
    if not m:
        m = re.search(r"L\s*-\s*T\s*-\s*P\s*[:\-]?\s*(\d)\s*-\s*(\d)\s*-\s*(\d)", head, re.I)
    if m:
        result["lecture_hours"], result["tutorial_hours"], result["practical_hours"] = (int(x) for x in m.groups())

    section = None
    current_unit = None
    unit_body = []
    objective_lines = []

    def flush_unit():
        nonlocal current_unit, unit_body
        if current_unit is not None:
            body = " ".join(unit_body)
            body, hours = _extract_hours(body) if current_unit["hours"] is None else (body, None)
            if hours and current_unit["hours"] is None:
                current_unit["hours"] = hours
            current_unit["topics"] = split_topics(body)
            if not current_unit["title"] and current_unit["topics"]:
                current_unit["title"] = current_unit["topics"].pop(0)
            result["units"].append(current_unit)
        current_unit, unit_body = None, []

    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        unit_match = UNIT_RE.match(line)
        new_section = None if unit_match else _section_of(line)
        if unit_match and _to_int(unit_match.group("num")):
            flush_unit()
            section = "units"
            rest = unit_match.group("rest")
            title, hours = _extract_hours(rest)
            # "Unit 1: Title: topic, topic" -> split at the first colon when content follows
            body = ""
            if ":" in title:
                title, body = title.split(":", 1)
            elif len(title) > 90 and "," in title:
                title, body = "", title
            current_unit = {"number": _to_int(unit_match.group("num")), "title": _clean(title), "hours": hours, "topics": []}
            if body.strip():
                unit_body.append(body)
            continue
        if new_section:
            flush_unit()
            section = new_section
            remainder = _clean(SECTION_PATTERNS[new_section].sub("", line, count=1))
            if remainder and new_section == "prerequisites":
                result["prerequisites"] = remainder
            continue

        co = CO_RE.match(line)
        if co and section in (None, "outcomes", "objectives", "contents", "assessment"):
            code = re.sub(r"[^0-9]", "", co.group("code"))
            result["outcomes"].append({"code": f"CO{int(code)}", "description": _clean(co.group("text"))})
            section = "outcomes"
            continue

        if section == "units" and current_unit is not None:
            unit_body.append(line)
        elif section == "outcomes":
            nm = NUMBERED_RE.match(line)
            if nm:
                result["outcomes"].append({"code": f"CO{len(result['outcomes']) + 1}", "description": _clean(nm.group(1))})
            elif result["outcomes"] and not re.search(r"able to|students will|on successful", line, re.I):
                result["outcomes"][-1]["description"] += " " + _clean(line)
        elif section == "objectives":
            objective_lines.append(_clean(line))
        elif section in ("textbooks", "references", "online", "experiments"):
            nm = NUMBERED_RE.match(line)
            items = result[section]
            if nm:
                items.append(_clean(nm.group(1)))
            elif items and not line[:1].isupper() or (items and len(line) < 25):
                items[-1] += " " + _clean(line)
            else:
                items.append(_clean(line))
        elif section == "prerequisites":
            result["prerequisites"] = (result["prerequisites"] + " " + _clean(line)).strip()
    flush_unit()

    # Units numbered twice (e.g. table of contents then detail) - keep the richer copy.
    by_number = {}
    for unit in result["units"]:
        prev = by_number.get(unit["number"])
        if prev is None or len(unit["topics"]) > len(prev["topics"]):
            by_number[unit["number"]] = unit
    result["units"] = [by_number[n] for n in sorted(by_number)]
    for unit in result["units"]:
        if not unit["title"]:
            unit["title"] = f"Unit {unit['number']}"
        if not unit["topics"]:
            unit["topics"] = [unit["title"]]

    # normalise outcome codes / de-duplicate
    seen, outcomes = set(), []
    for o in result["outcomes"]:
        if o["code"] not in seen and o["description"]:
            seen.add(o["code"])
            outcomes.append(o)
    result["outcomes"] = outcomes
    if not result["outcomes"] and objective_lines:
        result["outcomes"] = [
            {"code": f"CO{i + 1}", "description": re.sub(r"^to\s+", "", o, flags=re.I)}
            for i, o in enumerate(objective_lines[:6])
        ]
    result["experiments"] = [e for e in result["experiments"] if len(e) > 3]
    return result


# --- editable text round-trip ----------------------------------------------------

def to_editable(data):
    """Render parsed data into the text blocks shown on the review screen."""
    outcomes = "\n".join(f"{o['code']}: {o['description']}" for o in data.get("outcomes", []))
    unit_blocks = []
    for u in data.get("units", []):
        hours = f" | {u['hours']}" if u.get("hours") else ""
        block = [f"Unit {u['number']}: {u['title']}{hours}"]
        block += [f"- {t}" for t in u.get("topics", [])]
        unit_blocks.append("\n".join(block))
    return {
        "outcomes_text": outcomes,
        "units_text": "\n\n".join(unit_blocks),
        "textbooks_text": "\n".join(data.get("textbooks", [])),
        "references_text": "\n".join(data.get("references", []) + data.get("online", [])),
        "experiments_text": "\n".join(data.get("experiments", [])),
    }


def parse_outcomes_text(text):
    outcomes = []
    for line in (text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        m = CO_RE.match(line)
        if m:
            code = f"CO{int(re.sub(r'[^0-9]', '', m.group('code')))}"
            outcomes.append({"code": code, "description": _clean(m.group("text"))})
        else:
            outcomes.append({"code": f"CO{len(outcomes) + 1}", "description": _clean(line)})
    return outcomes


def parse_units_text(text):
    units = []
    for line in (text or "").splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        m = UNIT_RE.match(stripped)
        if m and _to_int(m.group("num")):
            rest = m.group("rest")
            hours = None
            if "|" in rest:
                rest, h = rest.rsplit("|", 1)
                hours = int(h.strip()) if h.strip().isdigit() else None
            else:
                rest, hours = _extract_hours(rest)
            units.append({"number": _to_int(m.group("num")), "title": _clean(rest) or f"Unit {m.group('num')}", "hours": hours, "topics": []})
        elif units:
            topic = _clean(stripped.lstrip("-•* "))
            if topic:
                units[-1]["topics"].append(topic)
    for i, u in enumerate(units):
        if not u["topics"]:
            u["topics"] = [u["title"]]
    return units


def parse_lines(text):
    return [_clean(NUMBERED_RE.sub(r"\1", l.strip())) for l in (text or "").splitlines() if l.strip()]
