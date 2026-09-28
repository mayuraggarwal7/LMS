from datetime import date, time, timedelta
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import User
from classroom import grading
from classroom.models import Assignment, Submission
from courses import syllabus
from courses.models import Course, Enrollment
from courses.setup import build_course
from delivery.models import AttendanceRecord, ClassSession, Holiday, TimetableSlot
from delivery.services import attendance_stats, sync_timetable
from exams.analytics import co_attainment, distribute_exam_total, exam_totals
from exams.models import ExamMark

SAMPLE = (Path(settings.BASE_DIR) / "samples" / "data_structures_syllabus.txt").read_text()


def make_course(teacher, **overrides):
    data = syllabus.parse_syllabus(SAMPLE)
    info = {"code": "CS201", "title": data["title"], "year": 2, "semester": 3, "credits": Decimal("4"),
            "lecture_hours": 3, "tutorial_hours": 0, "practical_hours": 2,
            "start_date": date(2026, 7, 6), "end_date": date(2026, 10, 9)}
    info.update(overrides)
    return build_course(teacher, data, info)


class SyllabusParserTests(TestCase):
    def test_parses_sample(self):
        data = syllabus.parse_syllabus(SAMPLE)
        self.assertEqual(data["code"], "CS201")
        self.assertEqual(data["title"], "Data Structures and Algorithms")
        self.assertEqual((data["lecture_hours"], data["tutorial_hours"], data["practical_hours"]), (3, 0, 2))
        self.assertEqual(len(data["outcomes"]), 5)
        self.assertEqual([u["hours"] for u in data["units"]], [6, 8, 8, 7, 7])
        self.assertIn("Binary search tree", data["units"][2]["topics"])
        self.assertEqual(len(data["experiments"]), 9)
        self.assertEqual(len(data["textbooks"]), 2)

    def test_module_layout_and_roman_numerals(self):
        text = """Course Code: EE105  Course Title: Basic Electrical Engineering
Module II - AC Circuits 10 hrs
Phasors; RLC series circuits; power factor
Module I: DC Circuits (8 Hrs)
Kirchhoff's laws, mesh and nodal analysis, superposition theorem
"""
        data = syllabus.parse_syllabus(text)
        self.assertEqual([u["number"] for u in data["units"]], [1, 2])
        self.assertEqual(data["units"][1]["hours"], 10)
        self.assertIn("Mesh and nodal analysis", data["units"][0]["topics"])

    def test_editable_round_trip(self):
        data = syllabus.parse_syllabus(SAMPLE)
        text = syllabus.to_editable(data)
        units = syllabus.parse_units_text(text["units_text"])
        self.assertEqual(units, [{"number": u["number"], "title": u["title"], "hours": u["hours"], "topics": u["topics"]} for u in data["units"]])
        self.assertEqual(syllabus.parse_outcomes_text(text["outcomes_text"]), data["outcomes"])

    def test_extract_text_from_docx(self):
        import io

        import docx

        document = docx.Document()
        document.add_paragraph("Unit 1: Sets (4 Hours)")
        document.add_paragraph("Sets, relations, functions")
        buf = io.BytesIO()
        document.save(buf)
        upload = SimpleUploadedFile("s.docx", buf.getvalue())
        text = syllabus.extract_text(upload)
        self.assertEqual(syllabus.parse_syllabus(text)["units"][0]["topics"], ["Sets", "Relations", "Functions"])


class CourseSetupTests(TestCase):
    def setUp(self):
        self.teacher = User.objects.create_user("t", password="pw123456", role=User.TEACHER)

    def test_build_course_generates_everything(self):
        course, summary = make_course(self.teacher)
        self.assertEqual(summary["units"], 5)
        self.assertEqual(course.plan_items.filter(session_type="lecture").count(), 36)  # = syllabus hours
        self.assertEqual(course.plan_items.filter(session_type="lab").count(), 9)
        self.assertTrue(all(i.outcomes.exists() for i in course.plan_items.all()))
        self.assertEqual(sum(c.weight for c in course.components.all()), 100)
        self.assertTrue(course.components.filter(kind="TW").exists())  # practical course gets term work
        self.assertGreater(course.questions.count(), 50)
        self.assertEqual(course.assignments.filter(kind="assignment").count(), 5)
        mse = course.exams.get(component__kind="MSE")
        ese = course.exams.get(component__kind="ESE")
        self.assertEqual(mse.paper_total, mse.total_marks)
        self.assertEqual(ese.paper_total, ese.total_marks)
        # ESE and MSE papers should not share questions
        self.assertFalse(set(mse.paper.values_list("question", flat=True)) & set(ese.paper.values_list("question", flat=True)))
        # assignment rubrics are private copies mapped to COs
        a = course.assignments.filter(kind="assignment").first()
        self.assertFalse(a.rubric.is_template)
        self.assertTrue(all(c.outcome_id for c in a.rubric.criteria.all()))

    def test_final_year_gets_project_and_different_pedagogy(self):
        course, _ = make_course(self.teacher, year=4)
        self.assertTrue(course.assignments.filter(kind="project").exists())
        pedagogies = set(course.plan_items.values_list("pedagogy", flat=True))
        self.assertTrue(any("Case study" in p or "Project-based" in p for p in pedagogies))


class TimetableSyncTests(TestCase):
    def setUp(self):
        self.teacher = User.objects.create_user("t", password="pw123456", role=User.TEACHER)
        self.course, _ = make_course(self.teacher, start_date=date(2026, 7, 6), end_date=date(2026, 7, 31))
        TimetableSlot.objects.create(course=self.course, weekday=0, start_time=time(10), end_time=time(11))
        TimetableSlot.objects.create(course=self.course, weekday=2, start_time=time(10), end_time=time(11))
        TimetableSlot.objects.create(course=self.course, weekday=3, start_time=time(14), end_time=time(16), session_type="lab")

    def test_sync_creates_sessions_and_maps_plan_in_order(self):
        Holiday.objects.create(course=self.course, date=date(2026, 7, 8))
        result = sync_timetable(self.course)
        lectures = list(self.course.sessions.filter(session_type="lecture"))
        self.assertEqual(len(lectures), 7)  # 4 Mondays + 4 Wednesdays - 1 holiday
        self.assertNotIn(date(2026, 7, 8), [s.date for s in lectures])
        seqs = [s.plan_item.sequence for s in lectures]
        self.assertEqual(seqs, sorted(seqs))
        self.assertEqual(lectures[0].plan_item.sequence, 1)
        labs = self.course.sessions.filter(session_type="lab")
        self.assertTrue(all(s.plan_item.session_type == "lab" for s in labs))
        self.assertGreater(result["unscheduled"], 0)

    def test_cancel_shifts_plan_and_completed_sessions_are_kept(self):
        sync_timetable(self.course)
        first, second = self.course.sessions.filter(session_type="lecture")[:2]
        first.status = ClassSession.COMPLETED
        first.save()
        planned_second = second.plan_item
        second.status = ClassSession.CANCELLED
        second.save()
        sync_timetable(self.course)
        third = self.course.sessions.filter(session_type="lecture", status="scheduled").first()
        self.assertEqual(third.plan_item, planned_second)
        first.refresh_from_db()
        self.assertEqual(first.plan_item.sequence, 1)
        # re-sync is idempotent
        before = list(self.course.sessions.values_list("pk", "plan_item"))
        sync_timetable(self.course)
        self.assertEqual(before, list(self.course.sessions.values_list("pk", "plan_item")))

    def test_attendance_stats(self):
        sync_timetable(self.course)
        student = User.objects.create_user("s", password="pw123456")
        Enrollment.objects.create(course=self.course, student=student)
        sessions = list(self.course.sessions.all()[:4])
        for i, s in enumerate(sessions):
            s.status = ClassSession.COMPLETED
            s.save()
            AttendanceRecord.objects.create(session=s, student=student, status="A" if i == 0 else "P")
        stats = attendance_stats(self.course)[student.pk]
        self.assertEqual((stats["present"], stats["total"], stats["percent"]), (3, 4, 75.0))
        self.assertFalse(stats["at_risk"])


class GradingTests(TestCase):
    def setUp(self):
        self.teacher = User.objects.create_user("t", password="pw123456", role=User.TEACHER)
        self.course, _ = make_course(self.teacher)
        self.student = User.objects.create_user("s", password="pw123456")
        Enrollment.objects.create(course=self.course, student=self.student)
        self.assignment = self.course.assignments.filter(kind="assignment").first()
        self.assignment.status = Assignment.PUBLISHED
        self.assignment.save()

    def test_rubric_to_score(self):
        sub = grading.get_or_create_submission(self.assignment, self.student)
        entries = {c.pk: {"level_id": c.levels.first().pk, "points": None} for c in self.assignment.rubric.criteria.all()}
        grading.apply_rubric_scores(sub, entries)
        self.assertEqual(sub.score, self.assignment.max_points)  # all "Excellent"
        self.assertEqual(sub.status, Submission.RETURNED)

    def test_score_to_rubric(self):
        sub = grading.get_or_create_submission(self.assignment, self.student)
        grading.set_total_score(sub, Decimal("5"))
        scores = list(sub.criterion_scores.all())
        self.assertEqual(len(scores), self.assignment.rubric.criteria.count())
        self.assertTrue(all(s.derived for s in scores))
        self.assertEqual(sum(s.points for s in scores), self.assignment.rubric.total_points / 2)
        # and back again: rubric recomputes to the same total
        self.assertEqual(grading.score_from_criteria(sub), Decimal("5.00"))

    def test_gradebook_weighting_and_min_pass(self):
        sub = grading.get_or_create_submission(self.assignment, self.student)
        grading.set_total_score(sub, Decimal("8"))  # 80% in ISE-1
        ese = self.course.exams.get(component__kind="ESE")
        distribute_exam_total(ese, self.student, Decimal("30"))  # 30% < 40% minimum
        self.assertEqual(exam_totals(ese)[self.student.pk], Decimal("30"))
        row = grading.build_gradebook(self.course)["rows"][0]
        weights = {c["component"].name: c["percent"] for c in row["components"]}
        self.assertAlmostEqual(weights["ISE-1"], 80.0)
        self.assertAlmostEqual(weights["ESE"], 30.0)
        self.assertAlmostEqual(row["final_percent"], (80 * 10 + 30 * 50) / 60)
        self.assertEqual(row["grade"], "F")
        self.assertIn("ESE", row["failed_components"])

    def test_student_view_hides_unreleased_marks_and_draft_grades(self):
        sub = grading.get_or_create_submission(self.assignment, self.student)
        grading.set_total_score(sub, Decimal("8"), grader_return=False)  # draft grade
        mse = self.course.exams.get(component__kind="MSE")
        distribute_exam_total(mse, self.student, Decimal("20"))  # not released
        teacher_row = grading.build_gradebook(self.course)["rows"][0]
        student_row = grading.build_gradebook(self.course, [self.student], released_only=True)["rows"][0]
        self.assertIsNotNone(teacher_row["final_percent"])
        self.assertIsNone(student_row["final_percent"])
        mse.status = "marked"
        mse.save()
        student_row = grading.build_gradebook(self.course, [self.student], released_only=True)["rows"][0]
        self.assertAlmostEqual(student_row["final_percent"], 20 / 30 * 100)

    def test_co_attainment(self):
        ese = self.course.exams.get(component__kind="ESE")
        for eq in ese.paper.all():
            ExamMark.objects.create(exam_question=eq, student=self.student, marks=eq.marks)
        report = {r["outcome"].code: r for r in co_attainment(self.course)}
        assessed = [r for r in report.values() if r["ese"]["assessed"]]
        self.assertTrue(assessed)
        self.assertTrue(all(r["ese"]["level"] == 3 for r in assessed))


class WorkflowTests(TestCase):
    """End-to-end through the views: syllabus upload -> course -> timetable -> class -> grading."""

    def setUp(self):
        self.teacher = User.objects.create_user("t", password="pw123456", role=User.TEACHER, first_name="T")
        self.student = User.objects.create_user("s", password="pw123456", role=User.STUDENT, roll_no="R1")

    def test_full_teacher_and_student_flow(self):
        self.client.login(username="t", password="pw123456")
        upload = SimpleUploadedFile("syllabus.txt", SAMPLE.encode())
        r = self.client.post(reverse("course_new"), {"syllabus_file": upload, "year": 2, "semester": 3,
                                                     "start_date": "2026-07-06", "end_date": "2026-10-09"})
        self.assertRedirects(r, reverse("course_review"))
        r = self.client.get(reverse("course_review"))
        form = r.context["form"]
        post = {k: (v if v is not None else "") for k, v in form.initial.items()}
        post.update({"opt_plan": "on", "opt_scheme": "on", "opt_rubrics": "on", "opt_questions": "on",
                     "opt_assignments": "on", "opt_exams": "on", "credits": "4"})
        r = self.client.post(reverse("course_review"), post)
        course = Course.objects.get(code="CS201")
        self.assertRedirects(r, reverse("course_setup_summary", args=[course.pk]))
        self.assertEqual(course.plan_items.count(), 45)

        # timetable + sync through the UI
        url = reverse("timetable", args=[course.pk])
        self.client.post(url, {"action": "slot", "slot-weekday": 0, "slot-start_time": "10:00", "slot-end_time": "11:00", "slot-session_type": "lecture"})
        self.client.post(url, {"action": "sync"})
        self.assertTrue(course.sessions.exists())

        # student joins with the class code
        self.client.logout()
        self.client.login(username="s", password="pw123456")
        self.client.post(reverse("join_course"), {"code": course.join_code})
        self.assertTrue(Enrollment.objects.filter(course=course, student=self.student).exists())
        # students cannot open teacher-only pages
        self.assertEqual(self.client.get(url).status_code, 403)
        self.assertEqual(self.client.get(reverse("gradebook", args=[course.pk])).status_code, 403)

        # teacher opens check-in; student checks in with the code
        session = course.sessions.first()
        session.date = timezone.localdate()
        session.checkin_open = True
        session.save()
        self.client.post(reverse("student_checkin", args=[course.pk]), {"code": session.checkin_code})
        self.assertTrue(AttendanceRecord.objects.filter(session=session, student=self.student, self_checked_in=True).exists())

        # student turns in the published assignment
        a = course.assignments.filter(kind="assignment").first()
        a.status = Assignment.PUBLISHED
        a.save()
        self.client.post(reverse("assignment_detail", args=[course.pk, a.pk]), {"action": "turn_in", "text": "my answer"})
        sub = Submission.objects.get(assignment=a, student=self.student)
        self.assertEqual(sub.status, Submission.TURNED_IN)

        # teacher grades with rubric levels and marks the class delivered with attendance
        self.client.logout()
        self.client.login(username="t", password="pw123456")
        data = {"action": "return", "feedback": "nice"}
        for c in a.rubric.criteria.all():
            data[f"level-{c.pk}"] = c.levels.all()[1].pk  # "Good" = 75%
        self.client.post(reverse("grade_submission", args=[course.pk, a.pk, self.student.pk]), data)
        sub.refresh_from_db()
        self.assertEqual(sub.score, Decimal("7.50"))
        self.assertEqual(sub.status, Submission.RETURNED)

        r = self.client.post(reverse("session_detail", args=[course.pk, session.pk]),
                             {"action": "complete", "coverage": "full", "notes": "Covered ADTs", "notes_shared": "on",
                              f"att-{self.student.pk}": "P"})
        session.refresh_from_db()
        self.assertEqual(session.status, ClassSession.COMPLETED)

        # student sees grade + notes
        self.client.logout()
        self.client.login(username="s", password="pw123456")
        r = self.client.get(reverse("my_grades", args=[course.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(self.client.get(reverse("session_detail", args=[course.pk, session.pk])), "Covered ADTs")

    def test_exam_marks_entry_and_total_distribution(self):
        course, _ = make_course(self.teacher)
        Enrollment.objects.create(course=course, student=self.student)
        exam = course.exams.get(component__kind="MSE")
        self.client.login(username="t", password="pw123456")
        self.client.post(reverse("exam_marks", args=[course.pk, exam.pk]), {f"total-{self.student.pk}": "18", "release": "1"})
        self.assertEqual(exam_totals(exam)[self.student.pk], Decimal("18"))
        exam.refresh_from_db()
        self.assertEqual(exam.status, "marked")

    def test_quick_grade_csv_import(self):
        course, _ = make_course(self.teacher)
        Enrollment.objects.create(course=course, student=self.student)
        a = course.assignments.first()
        self.client.login(username="t", password="pw123456")
        csv_file = SimpleUploadedFile("g.csv", b"roll_no,score\nR1,6\nunknown,3\n")
        self.client.post(reverse("import_grades", args=[course.pk, a.pk]), {"file": csv_file})
        sub = Submission.objects.get(assignment=a, student=self.student)
        self.assertEqual(sub.score, Decimal("6.00"))
        self.assertTrue(sub.criterion_scores.exists())

    def test_ics_export(self):
        course, _ = make_course(self.teacher, start_date=timezone.localdate(), end_date=timezone.localdate() + timedelta(days=14))
        TimetableSlot.objects.create(course=course, weekday=0, start_time=time(9), end_time=time(10))
        sync_timetable(course)
        self.client.login(username="t", password="pw123456")
        r = self.client.get(reverse("calendar_ics", args=[course.pk]))
        self.assertContains(r, "BEGIN:VEVENT")
